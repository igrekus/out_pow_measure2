import os
import openpyxl
import pandas as pd

from collections import defaultdict
from openpyxl.chart import LineChart, Series, Reference
from openpyxl.chart.axis import ChartLines
from openpyxl.cell import Cell
from textwrap import dedent

from forgot_again.file import load_ast_if_exists, pprint_to_file, make_dirs, open_explorer_at
from forgot_again.string import now_timestamp

GIGA = 1_000_000_000
MEGA = 1_000_000
KILO = 1_000
MILLI = 1 / 1_000


class MeasureResult:
    device = 'vco'
    measurement_name = 'tune'
    path = 'xlsx'

    def __init__(self):
        self._secondaryParams = dict()
        self._raw = list()
        self._raw_x2 = list()
        self._raw_x3 = list()

        self._report = dict()

        self._processed = list()
        self._processed_x2 = list()
        self._processed_x3 = list()

        self.ready = False

        self.data1 = defaultdict(list)
        self.data2 = defaultdict(list)
        self.data3 = dict()
        self.data4 = dict()
        self.data5 = defaultdict(list)
        self.data6 = defaultdict(list)

        self.adjustment = load_ast_if_exists('adjust.ini', default=None)

    def __bool__(self):
        return self.ready

    def _process(self):
        harm_x2 = [list(d.values()) for d in self._raw_x2]
        harm_x2 = _find_deltas(harm_x2, self._processed)
        self.data3[1] = harm_x2
        self._processed_x2 = harm_x2

        harm_x3 = [list(d.values()) for d in self._raw_x3]
        harm_x3 = _find_deltas(harm_x3, self._processed)
        self.data4[1] = harm_x3
        self._processed_x3 = harm_x3

        self.ready = True

    def add_harmonics_measurement(self, x2, x3):
        self._raw_x2 = list(x2)
        self._raw_x3 = list(x3)

    def _process_point(self, data):
        u_src = data['u_src']
        u_control = data['u_control']

        f_tune = data['read_f'] / MEGA
        p_out = data['read_p']
        i_src = data['read_i'] / MILLI

        if self.adjustment is not None:
            point = self.adjustment[len(self._processed)]
            f_tune += point['f_tune']
            p_out += point['p_out']
            i_src += point['i_src']

        self._report = {
            'u_src': u_src,
            'u_control': u_control,
            'f_tune': f_tune,
            'p_out': p_out,
            'i_src': i_src,
        }

        self.data1[u_src].append([u_control, f_tune])
        self.data2[u_src].append([u_control, p_out])
        self.data5[u_src].append([u_control, i_src])

        if len(self._processed):
            # (f2 - f1) / (u2 - u1) * 100
            last_point = self._processed[-1]
            f1 = last_point['f_tune']
            f2 = f_tune
            u1 = last_point['u_control']
            u2 = u_control
            tune = (f2 - f1) / (u2 - u1) * 100
            self.data6[u_src].append([u_control, tune])
        self._processed.append({**self._report})

    def clear(self):
        self._secondaryParams.clear()
        self._raw.clear()
        self._raw_x2.clear()
        self._raw_x3.clear()

        self._report.clear()

        self._processed.clear()
        self._processed_x2.clear()
        self._processed_x3.clear()

        self.data1.clear()
        self.data2.clear()
        self.data3.clear()
        self.data4.clear()
        self.data5.clear()
        self.data6.clear()

        self.adjustment = load_ast_if_exists('adjust.ini', default=None)

        self.ready = False

    def set_secondary_params(self, params):
        self._secondaryParams = dict(**params.params)

    def add_point(self, data):
        self._raw.append(data)
        self._process_point(data)

    def save_adjustment_template(self):
        if self.adjustment is None:
            print('measured, saving template')
            self.adjustment = [{
                'u_src': p['u_src'],
                'u_control': p['u_control'],
                'f_tune': 0,
                'p_out': 0,
                'i_src': 0,
            } for p in self._processed]
            pprint_to_file('adjust.ini', self.adjustment)

    @property
    def report(self):
        return dedent("""        ???????????????? ??????????????:
        U??????, ??={u_src}
        U??????, ??={u_control}
        I??????, mA={i_src}

        ????????????????????:
        F??????, ??????={f_tune:0.3f}
        P??????, ??????={p_out:0.3f}
        """.format(**self._report))

    def export_excel(self):
        make_dirs(self.path)
        fn = self._secondaryParams.get('file_name', None) or f'{self.device}-{self.measurement_name}-{now_timestamp()}'
        file_name = f'./{self.path}/{fn}.xlsx'

        u_dr_1, u_dr_2, u_dr_3 = 0, 0, 0

        udrs = list({point['u_src']: 0 for point in self._processed}.keys())
        try:
            u_dr_1 = udrs[0]
            u_dr_2 = udrs[1]
            u_dr_3 = udrs[2]
        except IndexError:
            pass

        result_harmonics_x2 = [[u_dr_1] + row for row in self._processed_x2]
        result_harmonics_x3 = [[u_dr_1] + row for row in self._processed_x3]

        df_harm_2 = pd.DataFrame(result_harmonics_x2, columns=['U??????, ??', 'U??????, ??', 'P??????_2, ??????'])
        df_harm_3 = pd.DataFrame(result_harmonics_x3, columns=['U??????, ??', 'U??????, ??', 'P??????_3, ??????'])

        df = pd.DataFrame(self._processed)
        df.columns=['U??????, ??', 'U??????, ??', 'F??????, ??????', 'P??????, ??????', 'I??????, ????', ]

        df['fdiff'] = df.groupby('U??????, ??')['F??????, ??????'].diff().shift(-1)
        df['udiff'] = df.groupby('U??????, ??')['U??????, ??'].diff().shift(-1)
        df['S, ??????/??'] = df[df['fdiff'].notna()].apply(lambda row: (row['fdiff'] / row['udiff']) * 100, axis = 1)
        df = df.drop(['fdiff', 'udiff'], axis=1)

        df = pd.merge(df, df_harm_2, how='left', on=['U??????, ??', 'U??????, ??'])
        df = pd.merge(df, df_harm_3, how='left', on=['U??????, ??', 'U??????, ??'])

        df['P??????_2??????, ??????'] = df[df['P??????_2, ??????'].notna()].apply(lambda row: -(row['P??????, ??????'] - row['P??????_2, ??????']), axis=1)
        df['P??????_3??????, ??????'] = df[df['P??????_3, ??????'].notna()].apply(lambda row: -(row['P??????, ??????'] - row['P??????_3, ??????']), axis=1)
        df['S, ??????/??'] = df['S, ??????/??'].fillna(0)

        df_udr_1 = df[df['U??????, ??'] == u_dr_1]
        df_udr_2 = df[df['U??????, ??'] == u_dr_2]
        df_udr_3 = df[df['U??????, ??'] == u_dr_3]

        cols = len(df_udr_1.columns)
        rows = len(df_udr_1)

        udr_1_s = [df_udr_1.columns.values.tolist()] + df_udr_1.values.tolist()
        udr_2_s = [df_udr_2.columns.values.tolist()] + (df_udr_2.values.tolist() or [[''] * cols] * rows)
        udr_3_s = [df_udr_3.columns.values.tolist()] + (df_udr_3.values.tolist() or [[''] * cols] * rows)

        print(udr_1_s)
        print(udr_2_s)
        print(udr_3_s)

        wb = openpyxl.Workbook()
        ws = wb.active

        out = []
        for r1, r2, r3 in zip(udr_1_s, udr_2_s, udr_3_s):
            row = r1 + ['', ''] + r2 + ['', ''] + r3
            out.append(row)

        for row in out:
            ws.append(row)

        top_left_cell: Cell = ws.cell(row=rows + 4, column=2)
        dx = 9
        dy = 15

        _add_chart(
            ws=ws,
            xs=Reference(ws, range_string=f'{ws.title}!B2:B{rows + 1}'),
            ys=[
                Reference(ws, range_string=f'{ws.title}!C2:C{rows + 1}'),
                Reference(ws, range_string=f'{ws.title}!O2:O{rows + 1}'),
                Reference(ws, range_string=f'{ws.title}!AA2:AA{rows + 1}'),
            ],
            title='???????????????? ??????????????????????',
            loc=top_left_cell.offset(0, 0).coordinate,
            curve_labels=['U?????? = 4.7??', 'U?????? = 5.0??', 'U?????? = 5.3??'],
            ax_titles=['U??????, ??', 'F??????, ??????'],
        )

        _add_chart(
            ws=ws,
            xs=Reference(ws, range_string=f'{ws.title}!B2:B{rows + 1}'),
            ys=[
                Reference(ws, range_string=f'{ws.title}!D2:D{rows + 1}'),
                Reference(ws, range_string=f'{ws.title}!P2:P{rows + 1}'),
                Reference(ws, range_string=f'{ws.title}!AB2:AB{rows + 1}'),
            ],
            title='????????????????',
            loc=top_left_cell.offset(0, dx).coordinate,
            curve_labels=['U?????? = 4.7??', 'U?????? = 5.0??', 'U?????? = 5.3??'],
            ax_titles=['U??????, ??', 'P??????, ??????'],
        )

        _add_chart(
            ws=ws,
            xs=Reference(ws, range_string=f'{ws.title}!B2:B{rows + 1}'),
            ys=[
                Reference(ws, range_string=f'{ws.title}!I2:I{rows + 1}'),
            ],
            title='?????????????????????????? ?????????????? 2?? ??????????????????',
            loc=top_left_cell.offset(dy, 0).coordinate,
            curve_labels=['U?????? = 4.7??'],
            ax_titles=['U??????, ??', 'P?????? ??2, ??????'],
        )

        _add_chart(
            ws=ws,
            xs=Reference(ws, range_string=f'{ws.title}!B2:B{rows + 1}'),
            ys=[
                Reference(ws, range_string=f'{ws.title}!J2:J{rows + 1}'),
            ],
            title='?????????????????????????? ?????????????? 3?? ??????????????????',
            loc=top_left_cell.offset(dy, dx).coordinate,
            curve_labels=['U?????? = 4.7??'],
            ax_titles=['U??????, ??', 'P?????? ??3, ??????'],
        )

        _add_chart(
            ws=ws,
            xs=Reference(ws, range_string=f'{ws.title}!B2:B{rows + 1}'),
            ys=[
                Reference(ws, range_string=f'{ws.title}!E2:E{rows + 1}'),
            ],
            title='?????? ??????????????????????',
            loc=top_left_cell.offset(0, 2 * dx).coordinate,
            curve_labels=['U?????? = 4.7??'],
            ax_titles=['U??????, ??', 'I??????, ????'],
        )

        _add_chart(
            ws=ws,
            xs=Reference(ws, range_string=f'{ws.title}!B2:B{rows}'),
            ys=[
                Reference(ws, range_string=f'{ws.title}!F2:F{rows}'),
            ],
            title='????????????????????????????????',
            loc=top_left_cell.offset(dy, 2 * dx).coordinate,
            curve_labels=['U?????? = 4.7??'],
            ax_titles=['U??????, ??', 'S, ??????/??'],
        )

        wb.save(file_name)
        open_explorer_at(os.path.abspath(file_name))


def _add_chart(ws, xs, ys, title, loc, curve_labels=None, ax_titles=None):
    chart = LineChart()

    for y, label in zip(ys, curve_labels):
        ser = Series(y, title=label)
        chart.append(ser)

    chart.set_categories(xs)
    chart.title = title

    chart.x_axis.minorGridlines = ChartLines()
    chart.x_axis.tickLblPos = 'low'

    if ax_titles:
        chart.x_axis.title = ax_titles[0]
        chart.y_axis.title = ax_titles[1]
    # chart.x_axis.tickLblSkip = 3

    ws.add_chart(chart, loc)


def _find_deltas(harm, origin):
    return [[harm[0], -(main['p_out'] - harm[1])] for harm, main in zip(harm, origin)]
